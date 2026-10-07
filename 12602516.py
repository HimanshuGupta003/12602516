import openpyxl
from datetime import datetime, timedelta

START = datetime(2026, 8, 17)
END = datetime(2026, 9, 21)
FEELING = {"Excellent": 5, "Good": 4, "Neutral": 3, "Low": 2, "Stressed": 1}
SATISFACTION = {"Very Satisfied": 5, "Satisfied": 4, "Neutral": 3,
                "Unsatisfied": 2, "Very Unsatisfied": 1}
ENERGY = {"High": 3, "Medium": 2, "Low": 1}


def read_data(filename):
    wb = openpyxl.load_workbook(filename, data_only=True)
    ws = wb["Daily Log"]
    days, dates, invalid, excluded = [], [], [], []
    for row_number in range(6, ws.max_row + 1):
        row = []
        for column in range(1, 15):
            row.append(ws.cell(row_number, column).value)
        # Ignore unused template rows, including their formula cells.
        if row[:8] == [None] * 8 and row[10:] == [None] * 4:
            continue
        try:
            day = row[0]
            if not isinstance(day, datetime):
                raise ValueError("missing or invalid Excel date")
            day = datetime(day.year, day.month, day.day)
            if day < START or day > END:
                excluded.append(day.strftime("%d-%m-%Y") + ": outside project period")
                continue
            if day in dates:
                raise ValueError("duplicate date")
            values, total = [], 0
            for column in [1, 2, 3, 4, 5, 7]:
                value = row[column]
                if type(value) not in (int, float) or not 0 <= value <= 1440:
                    raise ValueError("missing or invalid duration")
                values.append(value)
                total += value
            if not 0 < total <= 1440:
                raise ValueError("tracked time must be above 0 and at most 1440")
            count = row[6]
            if type(count) not in (int, float) or not 0 <= count <= 1440 or count != int(count):
                raise ValueError("invalid class count")
            feeling = FEELING[str(row[10]).strip()]
            satisfaction = SATISFACTION[str(row[11]).strip()]
            energy = ENERGY[str(row[12]).strip()]
            # Index: sleep, fitness, study, coding, class, other, total,
            # free time, feeling score, satisfaction score, energy score.
            values += [total, 1440 - total, feeling, satisfaction, energy]
            days.append(values)
            dates.append(day)
        except (ValueError, KeyError) as error:
            invalid.append("Row " + str(row_number) + ": " + str(error))
    wb.close()
    missing = []
    day = START
    while day <= END:
        if day not in dates:
            missing.append(day.strftime("%d-%m-%Y"))
        day += timedelta(days=1)
    return days, missing, invalid, excluded


def average(days, column):
    if len(days) == 0:
        raise ValueError("No valid days to analyse")
    total = 0
    for day in days:
        total += day[column]
    return total / len(days)


def activity_indices(days):
    result = {}
    result["TPI"] = average(days, 3)
    result["AAI"] = average(days, 2) + average(days, 4)
    result["PhAI"] = average(days, 1)
    result["SRI"] = average(days, 0)
    result["ABI"] = average(days, 7)
    result["TUI"] = average(days, 6)
    result["EI"] = (average(days, 8) + average(days, 9) + average(days, 10)) / 3
    result["DCI"] = len(days) / ((END - START).days + 1) * 100
    # Use the report template weights. DCI has a weight of 0.05.
    result["PAI"] = (0.15 * result["TPI"] + 0.20 * result["AAI"]
                     + 0.15 * result["PhAI"] + 0.20 * result["SRI"]
                     + 0.15 * result["TUI"] + 0.10 * result["EI"]
                     + 0.05 * result["DCI"])
    return result


def correlation(days, x_column, y_column):
    # Pearson r, calculated with loops and basic arithmetic.
    if len(days) < 2:
        return None
    x_mean = average(days, x_column)
    y_mean = average(days, y_column)
    xy, xx, yy = 0, 0, 0
    for day in days:
        x = day[x_column] - x_mean
        y = day[y_column] - y_mean
        xy += x * y
        xx += x * x
        yy += y * y
    if xx == 0 or yy == 0:
        return None
    return xy / (xx * yy) ** 0.5


def main():
    try:
        days, missing, invalid, excluded = read_data("12602516.xlsx")
        indices = activity_indices(days)
        lines = ["MY DATA, MY STORY - HIMANSHU GUPTA (12602516)",
                 "MCA / D1P2635", "Period: 17 August to 21 September 2026",
                 "Expected days: " + str((END - START).days + 1),
                 "Valid days: " + str(len(days)), "Missing days: " + str(len(missing)),
                 "Invalid records: " + str(len(invalid)),
                 "Outside-period records: " + str(len(excluded)), ""]
        names = ["Sleep", "Fitness", "Study", "Coding", "Class", "Other", "Tracked", "Free"]
        for column in range(8):
            lines.append(names[column] + ": " + format(average(days, column), ".2f") + " min/day")
        lines.append("\nINDICES")
        for name in ["PAI", "TPI", "AAI", "PhAI", "SRI", "ABI", "TUI", "EI", "DCI"]:
            unit = " min/day"
            if name == "PAI":
                unit = " (prescribed weighted index)"
            elif name == "EI":
                unit = " / 5 (template scale)"
            elif name == "DCI":
                unit = " %"
            lines.append(name + ": " + format(indices[name], ".2f") + unit)
        lines.append("\nRELATIONSHIPS (Pearson r)")
        for label, x, y in [("Sleep-Energy", 0, 10), ("Study-Satisfaction", 2, 9),
                            ("Coding-Energy", 3, 10)]:
            value = correlation(days, x, y)
            text = "undefined (too few days or no variation)" if value is None else format(value, ".3f")
            lines.append(label + ": " + text)
        lines += ["\nMissing dates: " + ", ".join(missing)] + invalid + excluded
        output = "\n".join(lines)
        print(output)
        with open("12602516_results.txt", "w", encoding="utf-8") as file:
            file.write(output + "\n")
    except (OSError, KeyError, ValueError) as error:
        print("Error:", error)


if __name__ == "__main__":
    main()
